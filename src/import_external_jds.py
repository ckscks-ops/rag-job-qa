# -*- coding: utf-8 -*-
"""
import_external_jds.py —— 把开源数据集转成本项目的 JSON，并与现有语料合并

数据源
======
`data/external/job_posting_data.xlsx`（由 `fetch_external_jds.py` 下载）
    10,114 条 2025 年中文招聘 JD，来自 Boss直聘 / 智联 / 猎聘
    label 分五类：数据分析 2650 / 经济金融 2095 / 生物统计 1939 / 算法岗 1898 / 大模型相关 1532

为什么要它
==========
本机 RecruitOps 导出的语料去重后 83.6% 是算法岗，**数据分析岗只有 122 条**，
与求职方向（数据分析 / AI 应用开发）不匹配，任何"数据分析岗都在要求什么"的结论都不可靠。
这个数据集补上 2,650 条数据分析岗。

⚠️ 诚实说明
  1. 这个数据集的岗位**带薪资，基本都是社招**，而 RecruitOps 那批是 2027 届校招。
     所以合并后语料是**校招 + 社招混合**，`job_type` 字段区分二者，评测时可按需过滤。
  2. 该仓库**没有 LICENSE**，数据来自公开招聘平台的岗位描述（本身公开）。
     个人求职作品集可用，但**不要商用、不要再分发原始文件**，引用要注明来源。

用法
----
    python src/import_external_jds.py --stats-only    # 只看转换结果，不写文件
    python src/import_external_jds.py --merge         # 转换 + 与现有语料合并
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / 'data' / 'external' / 'job_posting_data.xlsx'
BASE_CORPUS = ROOT / 'data' / 'prod_jobs.dedup.json'
OUT_EXTERNAL = ROOT / 'data' / 'external_jobs.json'
OUT_MERGED = ROOT / 'data' / 'corpus.json'

MIN_CHARS = 80


def _clean(s) -> str:
    if s is None:
        return ''
    s = str(s).strip()
    if s.lower() in ('nan', 'none'):
        return ''
    return s


def _norm_key(text: str) -> str:
    """去重用的归一化键：去掉所有空白。"""
    return re.sub(r'\s+', '', text or '')


def convert(xlsx: Path) -> list[dict]:
    """
    读 xlsx 并转成本项目的 JSON 结构。

    ⚠️ 用 **openpyxl** 而不是 pandas：
        pandas 会带进来一大堆依赖（约 12 MB + 编译轮子），
        而这里只需要顺序读一个 sheet——openpyxl 只读模式够用且轻得多，
        这样别人 `pip install -r requirements.txt` 后就能跑通全流程。
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        print('[错误] 需要 openpyxl：pip install openpyxl')
        raise

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    try:
        header = [str(h).strip() if h is not None else '' for h in next(it)]
    except StopIteration:
        print('  [错误] 表格是空的')
        return []
    idx = {name: i for i, name in enumerate(header)}
    print('  工作表: %s' % wb.sheetnames[0])
    print('  列    : %s' % header)

    def cell(row, name):
        i = idx.get(name)
        if i is None or i >= len(row):
            return None
        return row[i]

    jobs, dropped = [], Counter()
    n = 0
    for r in it:
        n += 1
        title = _clean(cell(r, 'jobname'))
        company = _clean(cell(r, 'company')) or '未知公司'
        city = _clean(cell(r, 'city'))
        desc = _clean(cell(r, 'description'))
        other = _clean(cell(r, 'other'))
        label = _clean(cell(r, 'label'))

        if not title:
            dropped['无岗位名'] += 1
            continue
        if len(desc) < MIN_CHARS:
            dropped['JD 过短'] += 1
            continue

        # 福利单独成段，让切块器能识别出「福利待遇」段
        raw = desc
        if other and len(other) >= 10:
            raw = desc.rstrip() + '\n福利待遇：' + other

        meta: dict = {'label': label, 'label_source': 'dataset'}
        for src, dst in (('minsalary', 'salary_min'), ('maxsalary', 'salary_max'),
                         ('meansalary', 'salary_mean')):
            v = cell(r, src)
            try:
                if v is not None and str(v).lower() != 'nan':
                    meta[dst] = round(float(v), 1)
            except (TypeError, ValueError):
                pass

        # ⚠️ doc_key 必须唯一（2026-10-09 踩坑）
        #    最初用 `ext|公司|岗位|城市`，结果同一公司同岗位同城市的不同帖子撞成一个文档，
        #    insert_document 的 ON CONFLICT 把它们静默合并，后面的重复插入又触发了
        #    contentless FTS 表的 DELETE 报错。正文哈希保证唯一且跨次运行稳定。
        h = hashlib.sha1(_norm_key(raw).encode('utf-8')).hexdigest()[:8]
        jobs.append({
            'doc_key': 'ext|%s|%s|%s|%s' % (company, title, city, h),
            'company': company,
            'title': title,
            'city': city,
            'job_type': '社招',                    # 带薪资，基本是社招
            'source': 'Job-posting-data(GitHub)',
            'raw_text': raw,
            **meta,
        })

    wb.close()
    print('  读到 %d 行' % n)
    if dropped:
        print('  丢弃: %s' % dict(dropped))
    return jobs


def dedup(jobs: list[dict], seen: set[str] | None = None) -> tuple[list[dict], int]:
    """按归一化正文去重；seen 可传入已有语料的键以实现**跨来源**去重。"""
    seen = set() if seen is None else seen
    out, dup = [], 0
    for j in jobs:
        k = _norm_key(j['raw_text'])
        if k in seen:
            dup += 1
            continue
        seen.add(k)
        out.append(j)
    return out, dup


def print_stats(jobs: list[dict], title: str) -> None:
    print()
    print('  === %s ===' % title)
    print('    条数          : %d' % len(jobs))
    lens = sorted(len(j['raw_text']) for j in jobs)
    if lens:
        print('    JD 长度       : min=%d 中位=%d p95=%d max=%d'
              % (lens[0], lens[len(lens) // 2], lens[int(len(lens) * .95)], lens[-1]))
    print('    公司(去重)    : %d' % len({j['company'] for j in jobs}))
    lab = Counter(j.get('label', '(无)') for j in jobs)
    if lab and list(lab) != ['(无)']:
        print('    方向标签      :')
        for k, v in lab.most_common():
            print('      %-14s %5d' % (k, v))
    print('    城市 Top5     : %s'
          % ', '.join('%s(%d)' % kv for kv in Counter(
              (j.get('city') or '(空)') for j in jobs).most_common(5)))


def main() -> int:
    ap = argparse.ArgumentParser(description='导入开源招聘数据集')
    ap.add_argument('--xlsx', default=str(XLSX))
    ap.add_argument('--stats-only', action='store_true')
    ap.add_argument('--merge', action='store_true',
                    help='与 data/prod_jobs.dedup.json 合并成 data/corpus.json')
    args = ap.parse_args()

    xlsx = Path(args.xlsx)
    if not xlsx.is_file():
        print('[错误] 找不到 %s，先运行 src/fetch_external_jds.py' % xlsx)
        return 1

    print('=' * 76)
    print('导入开源招聘 JD 数据集')
    print('=' * 76)
    print('  来源: %s' % xlsx)

    jobs = convert(xlsx)
    jobs_unique, dup_in = dedup(jobs)
    print('  数据集内部去重: %d -> %d（去掉重复 %d）' % (len(jobs), len(jobs_unique), dup_in))
    print_stats(jobs_unique, '导入结果')

    if args.stats_only:
        return 0

    OUT_EXTERNAL.write_text(json.dumps(jobs_unique, ensure_ascii=False, indent=1),
                            encoding='utf-8')
    print()
    print('  已写出: %s (%.2f MB)'
          % (OUT_EXTERNAL.relative_to(ROOT), OUT_EXTERNAL.stat().st_size / 1048576))

    if not args.merge:
        return 0

    # ---- 与现有语料合并 ----
    # ⚠️ 别人克隆本仓库时**没有** data/prod_jobs.dedup.json
    #    （那是我从本机 RecruitOps 的 pg_dump 备份里导出的，依赖本机安装）。
    #    这种情况不能直接 return —— 否则 corpus.json 永远不会生成，
    #    后面的 ingest 就没输入了。直接用开源数据集当语料。
    if not BASE_CORPUS.is_file():
        print()
        print('  [提示] 没有找到 %s' % BASE_CORPUS.name)
        print('         这一步依赖本机安装了 RecruitOps（别人的机器上不会有），')
        print('         所以直接用开源数据集作为语料。')
        print('         作者的完整语料是 11,334 条；克隆后是 %d 条。' % len(jobs_unique))
        OUT_MERGED.write_text(json.dumps(jobs_unique, ensure_ascii=False, indent=1),
                              encoding='utf-8')
        print()
        print('  ✅ 语料: %s (%.2f MB)'
              % (OUT_MERGED.relative_to(ROOT), OUT_MERGED.stat().st_size / 1048576))
        print_stats(jobs_unique, '语料（仅开源数据集）')
        return 0

    base = json.loads(BASE_CORPUS.read_text(encoding='utf-8'))
    print()
    print('  合并前：本机语料 %d 条 + 开源数据集 %d 条' % (len(base), len(jobs_unique)))

    # 现有语料优先（先占 key），开源数据集里重复的被丢掉
    seen = {_norm_key(j['raw_text']) for j in base}
    ext_new, cross_dup = dedup(jobs_unique, seen)
    merged = base + ext_new
    print('  跨来源去重: 开源数据集里 %d 条与现有语料重复，已丢弃' % cross_dup)

    # 现有语料内部可能仍有重复（prod_jobs.dedup.json 是按自己那批去重的）
    merged2, d2 = dedup(merged)
    if d2:
        print('  整体再查一遍: 又去掉 %d 条重复' % d2)
    merged = merged2

    OUT_MERGED.write_text(json.dumps(merged, ensure_ascii=False, indent=1),
                          encoding='utf-8')
    print()
    print('  ✅ 合并语料: %s (%.2f MB)'
          % (OUT_MERGED.relative_to(ROOT), OUT_MERGED.stat().st_size / 1048576))
    print_stats(merged, '合并后语料')

    src = Counter(j.get('source', '(未知)') for j in merged)
    print()
    print('    来源分布:')
    for k, v in src.most_common():
        print('      %-28s %5d' % (k, v))
    jt = Counter(j.get('job_type') or '(空)' for j in merged)
    print('    批次分布: %s' % dict(jt))

    # doc_key 唯一性自检（这是踩过坑的地方，必须每次检查）
    keys = [j['doc_key'] for j in merged]
    dup = [k for k, n in Counter(keys).items() if n > 1]
    print()
    if dup:
        print('    ❌ doc_key 重复 %d 个！例：%s' % (len(dup), dup[:3]))
        return 1
    print('    ✅ doc_key 全部唯一（%d 个）' % len(keys))
    return 0


if __name__ == '__main__':
    sys.exit(main())
